"""Read-only, source-linked rules from the supplied TSDR outsheet."""

from __future__ import annotations

from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
import math
import re


DEFAULT_OUTSHEET = Path(__file__).with_name("p1281_1_opt97_X81A_26WW41.3.xlsx")


def _text(value) -> str:
    return "" if value is None else str(value).strip()


@dataclass(frozen=True)
class Rule:
    label: str
    description: str
    drawn: str
    sheet: str
    row: int
    section: str
    subsection: str
    origin: str
    slv: str = ""
    sll: str = ""
    operator: str = ""

    @property
    def numeric_um(self) -> float | None:
        try:
            value = float(self.drawn)
            return value if math.isfinite(value) else None
        except ValueError:
            return None

    @property
    def gating(self) -> str:
        if "_NOTE" in self.label:
            return "Reference note"
        return "Tape-in gating" if self.label.startswith("TG") else "Approval required to waive" if self.label.startswith("TC") else "Reference"

    @property
    def reference(self) -> str:
        return f"{self.sheet}!D{self.row}" if self.sheet == "out_sheet" else f"{self.sheet}!B{self.row}"

    @property
    def plain(self) -> str:
        from dr_plain import simplify

        return simplify(self)

    @property
    def plain_gating(self) -> str:
        from dr_plain import plain_gating

        return plain_gating(self)

    @property
    def plain_section(self) -> str:
        from dr_plain import section_summary

        return section_summary(self.section)


@dataclass(frozen=True)
class RuleBook:
    source: Path
    process_tag: str
    option: str
    rules: tuple[Rule, ...]

    def search(self, query: str = "") -> list[Rule]:
        terms = re.findall(r"[a-z0-9]+", query.casefold())
        matches = [rule for rule in self.rules if all(term in " ".join((
            rule.label, rule.description, rule.section, rule.subsection, rule.origin, rule.plain)).casefold() for term in terms)]
        return sorted(matches, key=lambda rule: (rule.label.casefold() != query.strip().casefold(), rule.sheet, rule.row))

    def by_label(self, label: str) -> Rule | None:
        return next((rule for rule in self.rules if rule.label.casefold() == label.casefold()), None)

    @property
    def scope(self) -> str:
        return f"{self.source.name} ({self.process_tag}, {self.option}); X81A reference, not certified X81Z signoff"


def load_rules(path: Path = DEFAULT_OUTSHEET) -> RuleBook:
    path = Path(path).resolve()
    stamp = path.stat()
    return _load_rules(str(path), stamp.st_mtime_ns, stamp.st_size)


@lru_cache(maxsize=2)
def _load_rules(path: str, modified: int, size: int) -> RuleBook:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        metadata = list(workbook["p1281_ta_opt_info"].iter_rows(values_only=True))
        process, option = map(_text, metadata[1][:2])
        rules = []
        for sheet_name in ("out_sheet", "TC_SDR"):
            section = subsection = origin = ""
            for number, cells in enumerate(workbook[sheet_name].iter_rows(values_only=True), 1):
                values = [_text(value) for value in cells]
                if number == 1:
                    continue
                if sheet_name == "out_sheet":
                    kind, description, drawn, label, slv, sll = values[:6]
                    operator = ""
                else:
                    kind, label, description, operator, drawn = values[:5]
                    slv = sll = ""
                if kind == "header":
                    section, subsection = description, ""
                    origin = sll if sheet_name == "out_sheet" else ""
                elif kind == "subheader":
                    subsection = description
                elif label:
                    rules.append(Rule(label, description, drawn, sheet_name, number, section,
                                      subsection, origin, slv, sll, operator))
        return RuleBook(Path(path), process, option, tuple(rules))
    finally:
        workbook.close()


def source_status() -> str:
    try:
        book = load_rules()
        return f"TSDR source: {book.scope}; {len(book.rules)} labeled rules/notes indexed"
    except (OSError, ValueError, KeyError, ImportError) as exc:
        return f"TSDR outsheet unavailable: {exc}; no workbook-backed checks performed"


@dataclass(frozen=True)
class RuleReview:
    findings: tuple[str, ...] = ()
    checks: tuple[str, ...] = ()
    skipped: tuple[str, ...] = ()
    plain_findings: tuple[str, ...] = ()

    @property
    def summary(self) -> str:
        return (f"DR review INCOMPLETE: {len(self.checks)} evaluated checks/groups, "
                f"{len(self.findings)} finding(s). "
                "Selected unit-cell spacing/geometry checked; full enclosure, TSDR IDs and process applicability are not verified; not signoff.")

    @property
    def plain_summary(self) -> str:
        found = f"{len(self.findings)} possible problem(s) found" if self.findings else "no problems found"
        return (f"Design-rule check (partial): {len(self.checks)} checks ran, {found}. "
                "Only widths/spaces visible in the unit-cell preview are checked - not a signoff.")

    def display_findings(self, details: bool = False) -> tuple[str, ...]:
        return self.findings if details or not self.plain_findings else self.plain_findings

    def display_summary(self, details: bool = False) -> str:
        return self.summary if details else self.plain_summary


def review_rules(row: dict[str, str], fields: tuple[str, ...] | None = None) -> RuleReview:
    """Review modeled dimensions and explicitly account for unevaluated rules."""
    from ScalarCalculation import is_miller_row, segments, strip_shift
    from dr_plain import plain_check
    from scalar_snapshot import _resize_shift_nm

    pattern_fields = ("VcrPatternUnit", "VcrPatternGlobal", "VcrnubPatternUnit", "VcrnubPatternGlobal",
                      "VcrPattern", "VcrnubPattern")
    size_fields = ("PreviewVcrWidthNm", "PreviewVcrnubWidthNm", "PreviewVcrnubLengthNm", "CellHeight", "PolyPitch")
    geometry_fields = ("FTICdMisregSkew", "FTITracksUnit", "PolyCdSkew", "PolyPlugsUnit", "TCNCdRegSkew",
                       "TcnPlugsUnit", "VcxPatternUnit", "M0CutPatternUnit", "PreviewM0WidthNm",
                       "PreviewPlyWidthNm", "PreviewTcnWidthNm", "UnitCellSizeOgd", "UnitCellSizePgd")
    relevant = pattern_fields + size_fields + geometry_fields
    if fields is not None and not any(field in relevant for field in fields):
        return RuleReview(skipped=("Selected fields have no workbook-backed evaluator.",))
    try:
        book = load_rules()
    except (OSError, ValueError, KeyError, ImportError) as exc:
        return RuleReview(skipped=(f"Outsheet unavailable: {exc}",))
    if row.get("PolyPitch", "pp44").strip().lower() not in ("", "pp44"):
        return RuleReview(skipped=(f"PolyPitch {row.get('PolyPitch')!r}: only pp44 dimensional checks are implemented.",))
    match = re.fullmatch(r"(?:ch)?(110|132|165)", row.get("CellHeight", "").strip(), re.IGNORECASE)
    height = int(match[1]) if match else None
    if height not in (110, 132, 165):
        return RuleReview(skipped=(f"CellHeight {row.get('CellHeight')!r} is missing or unsupported.",))
    issues = []
    plain = {}
    checks = []
    skipped = []

    def dimension(field: str, default: float, location: str) -> float | None:
        try:
            value = float(row.get(field) or default)
            if not math.isfinite(value) or value <= 0:
                raise ValueError
            return value
        except (TypeError, ValueError):
            issues.append(f"Invalid dimension: {location} {field}={row.get(field)!r}; expected a positive finite nm value.")
            return None

    def check(value: float, label: str, operator: str, location: str, condition: str) -> None:
        rule = book.by_label(label)
        if rule is None or rule.numeric_um is None:
            skipped.append(f"{location}: {label} is missing or has no numeric limit in the outsheet.")
            return
        if not math.isfinite(value):
            issues.append(f"Invalid dimension: {location} is not finite.")
            return
        limit = rule.numeric_um * 1000
        checks.append(f"{label}: {location} {value:g} nm {operator} {limit:g} nm under {condition} (ID applicability unverified)")
        failed = value < limit - 1e-6 if operator == ">=" else value > limit + 1e-6 if operator == "<=" else abs(value - limit) > 1e-6
        if failed:
            issues.append(f"Outsheet advisory [{label}, {rule.reference}, X81A/opt97]: {location} "
                          f"{value:g} nm; expected {operator} {limit:g} nm under {condition}. "
                          "ID enclosure/applicability not verified; not signoff.")
            plain[issues[-1]] = plain_check(label, location, value, operator, limit)

    for field in pattern_fields:
        if fields is not None and field not in fields and not any(name in size_fields for name in fields):
            continue
        nub = field.startswith("Vcrnub")
        for index, segment in enumerate(segments(row.get(field, "")), 1):
            location = f"{field} segment {index}"
            parts = strip_shift(segment).split(",")
            if len(parts) < 2 or not parts[0].strip() or not parts[1].strip():
                issues.append(f"Invalid pattern: {location} needs X and Y axes; not evaluated.")
                continue
            if "|" in segment:
                try:
                    vector = [float(value) for value in segment.partition("|")[2].partition("$")[0].split(",")]
                    if not 1 <= len(vector) <= 4 or not all(math.isfinite(value) for value in vector):
                        raise ValueError
                except ValueError:
                    issues.append(f"Invalid skew: {location} needs 1..4 finite numeric values; not evaluated.")
                    continue
            cd_x, cd_y, reg_x, reg_y = _resize_shift_nm(segment)
            if not nub:
                base_width = dimension("PreviewVcrWidthNm", 12, location)
                if base_width is None:
                    continue
                width = base_width + cd_y
                if cd_x or cd_y or not (reg_x or reg_y):
                    check(width, "TGVCR_1010", ">=", location + " PGD width", "vcr:TSDRActiveCDID")
                    check(width, "TGVCR_1011", "<=", location + " PGD width", "vcr:TSDRActiveCDID")
                if reg_x or reg_y:
                    check(width, "TGVCR_1041", "=", location + " PGD width", "vcr:TSDRActiveMisregID")
                    check(abs(reg_y), "TGVCR_1045", "<=", location + " PGD displacement", "vcr:TSDRActiveMisregID")
                if reg_x:
                    skipped.append(f"{location}: OGD VCR registration is not modeled.")
                continue
            single = len(parts) > 2 and parts[2].strip().lower() in ("top", "bottom")
            pft = field in ("VcrPattern", "VcrnubPattern")
            default_ogd = 44 if is_miller_row(row) or pft else 132
            width = dimension("PreviewVcrnubWidthNm", default_ogd, location)
            nominal_length = (8.4 if height == 110 else 13.4) if pft else 13.4 if is_miller_row(row) else 26.8 if single and parts[2].strip().lower() == "top" else 28
            length = dimension("PreviewVcrnubLengthNm", nominal_length, location)
            if width is None or length is None:
                continue
            if default_ogd == 44:
                check(width + cd_x, "TGVCR_6017", ">=", location + " OGD width", "vcrnub:TSDRActiveCDID (1PP nub)")
                check(width + cd_x, "TGVCR_6018", "<=", location + " OGD width", "vcrnub:TSDRActiveCDID (1PP nub)")
            else:
                skipped.append(f"{location}: NoSpine/3PP nub OGD and spacing applicability are not evaluated.")
            if single:
                check(length + cd_y, "TGVCR_6011", ">=", location + " PGD length", "vcrnub:TSDRActiveCDID (single-sided)")
                check(length + cd_y, "TGVCR_6012", "<=", location + " PGD length", "vcrnub:TSDRActiveCDID (single-sided)")
            else:
                skipped.append(f"{location}: double-sided nub PGD geometry is not modeled.")
            if reg_x:
                check(abs(reg_x), "TGVCR_7045", "<=", location + " OGD displacement", "vcrnub:TSDRActiveMisregID")
            if reg_y:
                skipped.append(f"{location}: PGD nub registration and CRL alignment are not evaluated.")
    if fields is None or any(field in geometry_fields + ("CellHeight", "PolyPitch") for field in fields):
        from dr_geometry import review_geometry

        try:
            geometric = review_geometry(row)
            checks.extend(geometric.checks)
            skipped.extend(geometric.skipped)
            for finding in geometric.findings:
                issues.append(f"{finding.label} [{finding.reference or 'preview geometry'}, {finding.layer}]: "
                              f"{finding.message}; bounds nm {finding.bounds}")
                plain[issues[-1]] = finding.plain or finding.message
        except (ValueError, TypeError, KeyError, OverflowError) as exc:
            skipped.append(f"Geometric review unavailable: {exc}")
    if not checks:
        skipped.append("No supported dimensional comparisons were performed.")
    issues = list(dict.fromkeys(issues))
    return RuleReview(tuple(issues), tuple(checks), tuple(dict.fromkeys(skipped)),
                      tuple(plain.get(issue, issue) for issue in issues))


def rule_advisories(row: dict[str, str], fields: tuple[str, ...] | None = None) -> list[str]:
    return list(review_rules(row, fields).findings)