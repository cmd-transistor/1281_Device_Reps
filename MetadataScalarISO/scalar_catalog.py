"""PowerPoint catalog of the ISO scalar derivations (python-pptx).

Built from ``ISO_Scalar_Variants.csv`` and the snapshot PNGs written by
``ScalarCalculation.py --snapshots``:

1. title, executive summary (purpose, input, method, output, libraries,
   run command, caveats) and coverage (families / variant types) slides,
2. methodology + a concise summary table (family / type / fail-mode count),
3. one slide per variant: the exported layout snapshot with the structure
   name as caption and the ScalarFormula / notes underneath.

Used through ``ScalarCalculation.py --pptx FILE`` (full catalog) or
``--summary-pptx FILE`` (slides 1-2 only, no layout snapshots).
"""

from __future__ import annotations

import csv
import textwrap
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
SUMMARY_ROWS_PER_SLIDE = 42
SUMMARY_COLS = [("Family", 0.95), ("Type", 0.75), ("Variants", 0.7), ("Rows", 0.55), ("Cell heights", 1.05),
                ("FailModes", 0.85), ("Scalar", 0.95), ("Design patterns (StructureName suffixes)", 6.9)]

METHODOLOGY = [
    ("Scalar = 1 / FailModeCount", ""),
    ("PGD end-to-end (GATEGATE, EPIEPI, TCNTCN)",
     "FailModeCount = (# LO1<->HI1 tip-to-tip pairs per unit cell) x (# unit cells). A pair is a LO gate/TCN "
     "segment facing a HI segment end-to-end on the same column (adjacent diffusion tracks, or across an FTI). "
     "GATEGATE rows with vias removed (OC1/OC2/OC3) keep the SDR count from the declared LO1/HI1 tracks."),
    ("OGD contact-to-gate (CTG)",
     "FailModeCount = (# active PLY, i.e. gates with a VG) x (# diffusion tracks on which an active TCN -- a TCN "
     "with a VT -- runs beside the gate) x (# unit cells). One fail mode per gate per track: leakage happens on "
     "either side of the gate, the two sides are not added. Every variant (incl. VCR/VCRNUB types E/F) is "
     "counted from its own geometry."),
    ("OGD diffusion (DIFFOGD)",
     "FailModeCount = (# active FTI columns) x (# device diffusion tracks each FTI cuts, LO diffusion on one "
     "side / HI on the other) x (# unit cells)."),
    ("Diffusion chain (DIFFCHN)",
     "FailModeCount = # active VG gates in series (transistors in the current path, stacked gates included); "
     "multiplied by the number of chains when chains are stacked in an array (60x multi-instance rows)."),
    ("Unit cells", "ArraySizeOgd x ArraySizePgd from the buildsheet."),
    ("Source", "All counts are derived from the buildsheet row (FTITracksUnit, PolyPlugsUnit, TcnPlugsUnit, "
               "VcxPatternUnit, LO1/HI1ActiveTracks, DiffPatternUnit, array sizes). Via-track to diffusion-track "
               "mapping and the plug model were calibrated on the CTG132 chassis GDS bundled in ISO_LLM.bat."),
]


FAMILY_INFO = [
    ("CTG", "OGD contact-to-gate leakage: active gate (VG) beside an active TCN (VT) across the spacer",
     "A SDR reference (+PLYFSPCR/PLYFTCN CD skews); B/C/D gate vs TCN length combos; E VCR PGD skews; "
     "F VCRNUB OGD skews; G diff depop / block-EPI; H/I extra VG/VT in fail area; K/L density (DEN) variants"),
    ("GATEGATE", "PGD gate tip-to-tip leakage (LO1 gate end facing HI1 gate end)",
     "A 1TRK/2TRK gates; OC1/OC2/OC3 via removal (SDR count kept); 2TRK_4PP isolated sub-unit; "
     "C gate ETE through FTI (FTIML 1FC / FTISL HC) with OGD/PGD CD/REG skews"),
    ("EPIEPI", "PGD epi (TCN-landed diffusion) tip-to-tip leakage",
     "A 1TRK/2TRK reference; B epi-to-epi across FTI (OGD); C 4-TCN density; D VT closer to fail area; "
     "ESD/NES implant, PXL/NPXL/NDBI mask, GLK/LNK/PLG density variants"),
    ("TCNTCN", "PGD TCN tip-to-tip leakage",
     "A 1TRK/2TRK reference; B VCR/VCRNUB skews (+DEN); C 4-TCN density; D VT closer to fail area; "
     "TCP/TCN CD-REG skews; ESD/NES implant variants"),
    ("DIFFOGD", "OGD diffusion-to-diffusion leakage across an FTI cut",
     "A FTIML (1FC FTI cuts 2 device tracks) / FTISL (HC FTI cuts 1) with OGD/PGD CD/REG skews"),
    ("DIFFCHN", "Diffusion chain: transistors in series between LO1 and HI1",
     "A single gates (LEN3/6/12); B stacked gate pairs STKx1/STKx2; D serpentine 1FC/HC chains (~2xLEN gates); "
     "60x multi-instance arrays; PXL/FTI skews"),
]

HOW_IT_WORKS = [
    ("Purpose",
     "Fill the Scalar column (Scalar = 1 / fail-mode count) of the ISO E-test metadata buildsheet for every "
     "isolation structure and document the derivation (ScalarFormula) so each value can be audited."),
    ("Input",
     "ISO_Metadata_Input.csv -- the ISO buildsheet (one row per structure / test row). Columns used: "
     "StructureName, UnitCellSizeOgd/Pgd, ArraySizeOgd/Pgd, FTITracksUnit, PolyPlugsUnit, TcnPlugsUnit, "
     "VcxPatternUnit (vg/vt vias), LO1/HI1ActiveTracks, DiffPatternUnit, CellHeight."),
    ("Method",
     "1) Parse the buildsheet DSL (ranges a:b, strides a:b[N], rK end-relative tokens) into one Ogd x Pgd unit "
     "cell: FTI grid, gate and TCN segments (plug model), VG / VT positions, LO1 / HI1 nets.  "
     "2) Apply the family rule (CTG, GATEGATE/EPIEPI/TCNTCN, DIFFOGD, DIFFCHN) to count fail modes per unit.  "
     "3) Multiply by ArraySizeOgd x ArraySizePgd; Scalar = 1/N.  "
     "4) Group rows into geometry variants, render an annotated unit-cell snapshot per variant, build this deck."),
    ("Calibration",
     "Plug model (plug V merges diffusion tracks V-1/V) and via-track -> diffusion-track mapping were verified "
     "against the 12 CTG132 chassis GDS layouts bundled in ISO_LLM.bat; methodology per "
     "'x1281z - ISO E-test Scalar calculations suggestion.pptx'."),
    ("Output",
     "ISO_Metadata_Output.csv (input + Scalar + ScalarFormula) | ISO_Scalar_Review.csv (per-row derivation, "
     "intermediate counts, notes) | ISO_Scalar_Variants.csv (one line per geometry variant) | "
     "snapshots/*.png (optional, named after the structure) | ISO_Scalar_Catalog.pptx (this deck)."),
    ("Libraries",
     "Python 3.12 standard library (csv, re, argparse, dataclasses, pathlib) for the calculation; "
     "matplotlib for the snapshots; python-pptx for the catalog. No other dependencies."),
    ("Run",
     "python ScalarCalculation.py --pptx ISO_Scalar_Catalog.pptx   (add --keep-snapshots to retain PNGs; "
     "--snapshot NAME renders a single structure; --summary-pptx FILE writes this summary without layout slides)"),
    ("Scope / caveats",
     "Via-chain rows (VT_MV / VG_MV / VCR_MV) are out of scope and left blank. Snapshots are schematic "
     "(track -> diffusion-cell mapping is approximate). Data inconsistencies (e.g. LO1/HI1 tracks not matching "
     "the via tracks) are flagged in the Notes column rather than silently corrected."),
]


def _how_it_works_slide(prs: Presentation) -> None:
    s = _blank(prs)
    _text(s, Inches(0.5), Inches(0.25), Inches(12.3), Inches(0.6),
          ["Executive summary -- ISO scalar calculator"], size=24, bold_first=True)
    y = Inches(0.95)
    for head, body in HOW_IT_WORKS:
        lines = textwrap.wrap(body, 165)
        h = Inches(0.24) * len(lines) + Inches(0.08)
        _text(s, Inches(0.5), y, Inches(1.6), h, [head], size=11, bold_first=True)
        _text(s, Inches(2.1), y, Inches(10.8), h, lines, size=10)
        y += h + Inches(0.04)


def _coverage_slide(prs: Presentation, all_variants: list[dict[str, str]]) -> None:
    scored = [r for r in all_variants if r["FailModes"]]
    total_rows = sum(int(r["Rows"]) for r in all_variants)
    via_rows = sum(int(r["Rows"]) for r in all_variants if r["Family"] == "VIA_CHAIN")
    s = _blank(prs)
    _text(s, Inches(0.5), Inches(0.25), Inches(12.3), Inches(0.6),
          ["Coverage -- structure families and variant types"], size=24, bold_first=True)
    _text(s, Inches(0.5), Inches(0.85), Inches(12.3), Inches(0.4),
          [f"{total_rows} buildsheet rows: {total_rows - via_rows} scored across {len(scored)} geometry variants; "
           f"{via_rows} via-chain rows out of scope. Scalar spans "
           f"1/{max(int(r['FailModes']) for r in scored)} .. 1/{min(int(r['FailModes']) for r in scored)}."],
          size=11)
    cols = [("Family", 1.0), ("Rows", 0.6), ("Variants", 0.75), ("FailModes range", 1.3),
            ("Leakage path counted", 3.6), ("Variant types in the input", 5.45)]
    shape = s.shapes.add_table(len(FAMILY_INFO) + 1, len(cols), Inches(0.3), Inches(1.4), Inches(12.7),
                               Inches(0.75) * (len(FAMILY_INFO) + 1))
    tbl = shape.table
    for j, (name, w) in enumerate(cols):
        tbl.columns[j].width = Inches(w)
        c = tbl.cell(0, j)
        c.text = name
        c.text_frame.paragraphs[0].font.size = Pt(9)
        c.text_frame.paragraphs[0].font.bold = True
    for i, (fam, path, types) in enumerate(FAMILY_INFO, 1):
        fr = [r for r in scored if r["Family"] == fam]
        fms = sorted({int(r["FailModes"]) for r in fr})
        rng = f"{fms[0]} .. {fms[-1]}" if len(fms) > 1 else (str(fms[0]) if fms else "-")
        vals = [fam, str(sum(int(r["Rows"]) for r in fr)), str(len(fr)), rng, path, types]
        for j, v in enumerate(vals):
            c = tbl.cell(i, j)
            c.text = v
            for p in c.text_frame.paragraphs:
                p.font.size = Pt(8)
            c.margin_top = c.margin_bottom = Emu(18000)


def _blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _text(slide, left, top, width, height, lines, size=12, bold_first=False, color=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(size)
        p.font.bold = bold_first and i == 0
        if color is not None:
            p.font.color.rgb = color
    return box


def _title_slide(prs: Presentation, n_variants: int, n_rows: int) -> None:
    s = _blank(prs)
    _text(s, Inches(0.7), Inches(2.3), Inches(12), Inches(1.2),
          ["ISO E-test Scalar Catalog (1281 node)"], size=36, bold_first=True)
    _text(s, Inches(0.7), Inches(3.6), Inches(12), Inches(2),
          [f"{n_variants} structure variants covering {n_rows} buildsheet rows",
           "Per variant: annotated unit-cell layout, counted leakage paths and the ScalarFormula",
           "Generated by ScalarCalculation.py --pptx"], size=16)


def _methodology_slide(prs: Presentation) -> None:
    s = _blank(prs)
    _text(s, Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.7), ["Methodology"], size=26, bold_first=True)
    y = Inches(1.1)
    for head, body in METHODOLOGY:
        lines = [head] + textwrap.wrap(body, 150)
        h = Inches(0.32) * len(lines) + Inches(0.1)
        _text(s, Inches(0.5), y, Inches(12.3), h, lines, size=13, bold_first=True)
        y += h


def _summary_rows(variants: list[dict[str, str]]) -> list[list[str]]:
    """One line per (family, type, fail-mode count): variants, rows, heights, patterns."""
    groups: dict[tuple[str, str, int], dict] = {}
    for r in variants:
        fm = int(r["FailModes"] or 0)
        g = groups.setdefault((r["Family"], r["Type"], fm),
                              {"variants": 0, "rows": 0, "heights": set(), "patterns": [], "scalar": r["Scalar"]})
        g["variants"] += 1
        g["rows"] += int(r["Rows"])
        g["heights"].update(r["CellHeights"].split())
        pat = r["Variant"].split("-", 2)[-1]           # drop ``FAMILY-<MOS>-``
        if pat not in g["patterns"]:
            g["patterns"].append(pat)
    fam_order = {f: i for i, f in enumerate(["CTG", "GATEGATE", "EPIEPI", "TCNTCN", "DIFFOGD", "DIFFCHN"])}
    out: list[list[str]] = []
    for (fam, typ, fm), g in sorted(groups.items(), key=lambda kv: (fam_order.get(kv[0][0], 99), kv[0][1], kv[0][2])):
        pats = "; ".join(g["patterns"])
        out.append([fam, typ, str(g["variants"]), str(g["rows"]), " ".join(sorted(g["heights"])),
                    str(fm), g["scalar"], pats if len(pats) <= 150 else pats[:147] + "..."])
    return out


def _summary_slides(prs: Presentation, variants: list[dict[str, str]]) -> int:
    rows = _summary_rows(variants)
    pages = [rows[i:i + SUMMARY_ROWS_PER_SLIDE] for i in range(0, len(rows), SUMMARY_ROWS_PER_SLIDE)] or [[]]
    for pi, page in enumerate(pages, 1):
        s = _blank(prs)
        suffix = f" ({pi}/{len(pages)})" if len(pages) > 1 else ""
        _text(s, Inches(0.4), Inches(0.15), Inches(12.5), Inches(0.45),
              [f"Scalar summary by structure family / type{suffix}"], size=18, bold_first=True)
        _text(s, Inches(0.4), Inches(0.55), Inches(12.5), Inches(0.3),
              ["FailModes = fail-mode count of the whole structure (Scalar = 1/FailModes); full per-variant detail in "
               "ISO_Scalar_Variants.csv and on the following slides"], size=9, color=RGBColor(0x55, 0x55, 0x55))
        row_h = Inches(0.148)
        shape = s.shapes.add_table(len(page) + 1, len(SUMMARY_COLS), Inches(0.3), Inches(0.9),
                                   Inches(12.7), row_h * (len(page) + 1))
        tbl = shape.table
        for j, (name, w) in enumerate(SUMMARY_COLS):
            tbl.columns[j].width = Inches(w)
            cell = tbl.cell(0, j)
            cell.text = name
            cell.text_frame.paragraphs[0].font.size = Pt(8)
            cell.text_frame.paragraphs[0].font.bold = True
            cell.margin_top = cell.margin_bottom = Emu(0)
        for i, vals in enumerate(page, 1):
            tbl.rows[i].height = row_h
            for j, v in enumerate(vals):
                cell = tbl.cell(i, j)
                cell.text = v
                cell.text_frame.paragraphs[0].font.size = Pt(7)
                cell.margin_top = cell.margin_bottom = Emu(0)
                cell.margin_left = cell.margin_right = Emu(36000)
    return len(pages)


def _snapshot_slide(prs: Presentation, r: dict[str, str], png: Path) -> None:
    s = _blank(prs)
    _text(s, Inches(0.4), Inches(0.15), Inches(12.5), Inches(0.5),
          [r["ExampleStructure"]], size=18, bold_first=True)
    sub = (f"{r['VariantId']}  |  {r['Variant']}  |  {r['Rows']} rows  |  {r['CellHeights']}  |  "
           f"test rows: {r['TestRows']}")
    _text(s, Inches(0.4), Inches(0.6), Inches(12.5), Inches(0.35), [sub], size=10,
          color=RGBColor(0x55, 0x55, 0x55))
    pic_top, pic_h = Inches(0.95), Inches(4.7)
    pic = s.shapes.add_picture(str(png), Inches(0.4), pic_top, height=pic_h)
    if pic.width > SLIDE_W - Inches(0.8):
        pic.width, pic.height = SLIDE_W - Inches(0.8), int(pic.height * (SLIDE_W - Inches(0.8)) / pic.width)
    lines = [f"FailModes = {r['FailModes']}    Scalar = {r['Scalar']}"]
    lines += textwrap.wrap(r["Formula"], 175)
    if r["Notes"]:
        lines += textwrap.wrap("Notes: " + r["Notes"], 175)
    if r["StructureComments"]:
        lines += textwrap.wrap("Comment: " + r["StructureComments"], 175)
    _text(s, Inches(0.4), pic_top + pic.height + Inches(0.1), Inches(12.5), Inches(1.5), lines, size=10,
          bold_first=True)


def build_catalog(variants_csv: Path, snapshot_dir: Path | None, out_pptx: Path,
                  with_snapshots: bool = True) -> int:
    with open(variants_csv, newline="", encoding="utf-8-sig") as fh:
        all_variants = list(csv.DictReader(fh))
    variants = [r for r in all_variants if r.get("FailModes")]
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    _title_slide(prs, len(variants), sum(int(r["Rows"]) for r in variants))
    _how_it_works_slide(prs)
    _coverage_slide(prs, all_variants)
    _methodology_slide(prs)
    _summary_slides(prs, variants)
    if with_snapshots and snapshot_dir is not None:
        for r in variants:
            png = Path(snapshot_dir) / (r.get("Snapshot") or "")
            if r.get("Snapshot") and png.exists():
                _snapshot_slide(prs, r, png)
    prs.save(str(out_pptx))
    return len(prs.slides)
